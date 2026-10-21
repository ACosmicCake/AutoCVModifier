# app/cv_utils.py
import json
import logging

import docx

try:
    from pypdf import PdfReader
except ImportError:  # older installs
    from PyPDF2 import PdfReader

from .llm import call_llm, extract_json, strip_reasoning

logger = logging.getLogger(__name__)

SUMMARY_KEY = "SummaryOrObjective"


# --- CV file readers -------------------------------------------------------

def get_cv_from_text_file(filepath: str) -> str | None:
    """Reads CV content from a plain text file."""
    for encoding in ("utf-8", "utf-8-sig", "cp1252"):
        try:
            with open(filepath, "r", encoding=encoding) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
        except OSError as e:
            logger.error("Error reading text file %s: %s", filepath, e)
            return None
    logger.error("Could not decode text file %s", filepath)
    return None


def get_cv_from_json_file(filepath: str) -> dict | None:
    """Reads CV content from a JSON file and parses it into a dictionary."""
    try:
        with open(filepath, "r", encoding="utf-8") as f:
            return json.load(f)
    except (OSError, json.JSONDecodeError) as e:
        logger.error("Error reading JSON file %s: %s", filepath, e)
        return None


def get_cv_from_pdf_file(filepath: str) -> str | None:
    """Extracts text content from a PDF file. Returns None if the file can't be read."""
    try:
        reader = PdfReader(filepath)
        if reader.is_encrypted:
            reader.decrypt("")
        pages = [page.extract_text() or "" for page in reader.pages]
        text = "\n".join(p for p in pages if p.strip())
        if not text:
            logger.warning("No text extracted from PDF %s (it may be a scanned image).", filepath)
        return text
    except Exception as e:
        logger.error("Error processing PDF file %s: %s", filepath, e)
        return None


def get_cv_from_docx_file(filepath: str) -> str | None:
    """Extracts text from a DOCX file, including text inside tables (common in CV templates)."""
    try:
        document = docx.Document(filepath)
        lines = [para.text for para in document.paragraphs if para.text.strip()]
        for table in document.tables:
            for row in table.rows:
                seen = []
                for cell in row.cells:  # merged cells repeat, so de-duplicate per row
                    text = cell.text.strip()
                    if text and text not in seen:
                        seen.append(text)
                if seen:
                    lines.append(" | ".join(seen))
        return "\n".join(lines)
    except Exception as e:
        logger.error("Error processing DOCX file %s: %s", filepath, e)
        return None


def get_cv_content_from_file(filepath: str) -> str | None:
    """Extracts CV text from any supported file type."""
    ext = filepath.rsplit(".", 1)[-1].lower()
    if ext == "pdf":
        return get_cv_from_pdf_file(filepath)
    if ext == "docx":
        return get_cv_from_docx_file(filepath)
    if ext == "txt":
        return get_cv_from_text_file(filepath)
    if ext == "json":
        data = get_cv_from_json_file(filepath)
        return json.dumps(data, indent=2, ensure_ascii=False) if data else None
    return None


# --- CV template -----------------------------------------------------------

def load_cv_template(path: str, include_summary: bool) -> str:
    """
    Loads CV_format.json and returns it as a JSON string for the prompt.
    The summary section is dropped unless requested. Raises FileNotFoundError / ValueError.
    """
    with open(path, "r", encoding="utf-8") as f:
        raw = f.read()
    try:
        template = json.loads(raw)
    except json.JSONDecodeError as e:
        raise ValueError(f"CV format file is not valid JSON: {e}") from e

    cv = template.get("CV", template)
    if include_summary:
        cv.setdefault(SUMMARY_KEY, {"Statement": ""})
    else:
        cv.pop(SUMMARY_KEY, None)
    return json.dumps(template, indent=2)


def _is_blank(value) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return not value.strip()
    if isinstance(value, (list, tuple)):
        return all(_is_blank(v) for v in value)
    if isinstance(value, dict):
        return all(_is_blank(v) for v in value.values())
    return False


def normalize_cv_data(data: dict, include_summary: bool = True) -> dict:
    """
    Cleans up a model-produced CV: wraps it under a top-level "CV" key, removes
    template placeholder entries that were echoed back empty, and drops the summary
    if it wasn't requested.
    """
    if "CV" not in data or not isinstance(data.get("CV"), dict):
        data = {"CV": data}
    cv = data["CV"]

    for key, value in list(cv.items()):
        if isinstance(value, list):
            cleaned = []
            for item in value:
                if _is_blank(item):
                    continue
                if isinstance(item, dict):
                    item = {k: ([x for x in v if not _is_blank(x)] if isinstance(v, list) else v) for k, v in item.items()}
                cleaned.append(item)
            cv[key] = cleaned

    if not include_summary:
        cv.pop(SUMMARY_KEY, None)
    return data


# --- Prompts ---------------------------------------------------------------

SYSTEM_PROMPT = (
    "You are an expert technical recruiter and career coach. You write in plain, humble, natural English, "
    "and you never invent facts about a candidate: every employer, title, date, degree, certification and "
    "metric you write must be supported by the candidate's original CV."
)


def build_tailor_prompt(cv_content_str: str, job_description_text: str, cv_template_content_str: str,
                        include_summary: bool) -> str:
    summary_step = (
        "1.  **Summary**: Write a dense 3-4 line paragraph that maps the candidate's real qualifications to the "
        "role's primary responsibilities. Put it in `CV.SummaryOrObjective.Statement`.\n"
        if include_summary else
        "1.  **No Summary**: The structure intentionally has no summary section. Do not add one.\n"
    )
    return f"""
Your task is to tailor the candidate's CV to the job description below and return it as JSON.

**Goal: Dual Optimization**
1.  **ATS relevance**: Use the job description's important keywords wherever the candidate genuinely has that experience.
2.  **Human persuasion**: Present a clear, compelling story of value and impact. Use humble, straightforward, natural language. Avoid buzzwords such as "stakeholder", "visionary", "pioneer", "innovator", "synergy", "passionate".

**Phase 1: Analyse the job description**
Identify the must-have qualifications, the top 3-4 responsibilities, the business problem the role solves, and the high-value keywords (tools, technologies, methodologies).

**Phase 2: Populate the JSON**
{summary_step}2.  **Job Title**: Put a short job title for the target role (two key words, no special characters) in `CV.JobTitle`.
3.  **Experience**: Keep the candidate's real roles (employer, job title and dates exactly as given) in `CV.ProfessionalExperience`, most relevant and recent first. Rewrite each role's bullet points in `ResponsibilitiesAndAchievements` using the result-first C.A.R.L. method (Context, Action, Result, Learning): lead with the outcome, emphasise the work most relevant to this job, and use the job's terminology where it truthfully applies. Only use numbers that appear in, or can be directly derived from, the original CV.
4.  **Projects**: Feature the candidate's most relevant projects in `CV.Projects`, with a brief `Description` and result-oriented bullets in `KeyContributionsOrTechnologiesUsed`.
5.  **Skills**: Group the candidate's relevant skills into `CV.Skills` objects, each with a `SkillCategory` (e.g. "Programming Languages", "Cloud & DevOps", "Databases", "Frameworks & Libraries", "Methodologies") and a `Skill` array. Order by relevance to the job. Only list skills the candidate actually has.
6.  **Remaining sections**: Transfer `PersonalInformation`, `Education` and any other sections in the structure accurately from the original CV.

**Output rules**
* Return ONLY a single valid JSON object that follows the LABELING STRUCTURE exactly, starting with the top-level `CV` key. No markdown, no comments, no text before or after it.
* Use `""` or `[]` for fields that don't apply. Do not leave template placeholder entries with empty values.
* Be comprehensive but concise: the CV should fit on one to two pages.

--- BEGIN CV ---
{cv_content_str}
--- END CV ---

--- BEGIN JOB DESCRIPTION ---
{job_description_text}
--- END JOB DESCRIPTION ---

--- BEGIN LABELING STRUCTURE ---
{cv_template_content_str}
--- END LABELING STRUCTURE ---
"""


def process_cv_and_jd(cv_content_str: str, job_description_text: str, cv_template_content_str: str,
                      model: str | None = None, include_summary: bool = False) -> str | None:
    """
    Tailors the CV to the job description. Returns the tailored CV as a JSON string,
    or None if the input is missing or the model's output isn't valid JSON.
    Raises LLMError when the model call itself fails.
    """
    if not cv_content_str:
        logger.error("CV content is missing for processing.")
        return None
    if not job_description_text:
        logger.error("Job description is missing for processing.")
        return None

    prompt = build_tailor_prompt(cv_content_str, job_description_text, cv_template_content_str, include_summary)
    output = call_llm(prompt, model, system=SYSTEM_PROMPT, json_mode=True)

    try:
        data = extract_json(output)
    except ValueError as e:
        logger.error("Model output was not valid JSON: %s. First 500 chars: %s", e, output[:500])
        return None
    return json.dumps(normalize_cv_data(data, include_summary), indent=2, ensure_ascii=False)


def _as_json_text(cv_json) -> str:
    return cv_json if isinstance(cv_json, str) else json.dumps(cv_json, indent=2, ensure_ascii=False)


def generate_cover_letter(cv_json, job_description: str, model: str | None = None) -> str:
    """Generates a cover letter from the tailored CV. Raises LLMError on failure."""
    prompt = f"""
Write a compelling cover letter for this candidate, based on their tailored CV and the job description.

**Guiding principles**
*   **Narrative of value**: Don't just list skills. Show how the candidate's experience solves this employer's specific problems.
*   **Show, don't tell**: Instead of "experienced leader", describe a real situation from the CV where the candidate led a team to success.
*   **Humble, natural tone**: Straightforward and engaging. Avoid jargon such as "stakeholder", "visionary" or "pioneer".
*   **Truthful**: Only use facts from the CV.

**Structure**
1.  A strong opening that states the role and why the candidate is a fit.
2.  One or two body paragraphs connecting 2-3 of the most relevant achievements to the job's requirements.
3.  A confident close with a call to action, signed with the candidate's name.

Return only the letter text (no markdown, no placeholders like [Company Address] unless the information is unknown).

**Tailored CV:**
```json
{_as_json_text(cv_json)}
```

**Job Description:**
```
{job_description}
```
"""
    return strip_reasoning(call_llm(prompt, model, system=SYSTEM_PROMPT))


def answer_question(cv_json, job_description: str, question: str, model: str | None = None) -> str:
    """Answers one application question. Raises LLMError on failure."""
    prompt = f"""
Answer the following job application question on behalf of the candidate, using their CV and the job description.

**Guidelines**
*   Be direct and concise: answer the question in one short paragraph (or a few if the question demands it).
*   Use specific, real examples from the CV to support the answer.
*   Write in the first person, in a humble, natural tone. No bullet points or lists.
*   Return only the answer text, without repeating the question.

**Tailored CV:**
```json
{_as_json_text(cv_json)}
```

**Job Description:**
```
{job_description}
```

**Application Question:**
{question}
"""
    return strip_reasoning(call_llm(prompt, model, system=SYSTEM_PROMPT))
